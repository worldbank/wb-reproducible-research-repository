"""
Data inventory — for code-only mode (SKILL.md, Phase 0)

Run this on YOUR OWN machine. It lists every data file under a folder and
records file-level metadata only:
  filename, relative path, SHA256 hash, size, last-modified date, and
  (for .dta, .sav, .csv, .tsv, .xlsx, .parquet) variable/column names and
  Excel sheet names.
It never records data values, value labels, variable labels, or summary
statistics. Open data_inventory.csv and review it before sharing it with
the AI agent.

The empty columns at the end (source, url, access_date, license,
availability, notes) are for you to fill in for each dataset you did not
create yourself. They become the Data Availability Statement in the README.
Leave them blank for files your own code produces.

Hashing follows resources/generate_hash in this repository.

Only the standard library is required. Variable names are read with these
packages when installed (otherwise that file type gets hash and size only):
  .dta -> pandas    .sav -> pyreadstat    .xlsx -> openpyxl
  .xls -> pandas + xlrd    .parquet -> pyarrow
"""

# ------------------------------------------------------------------
# 1. CHANGE THESE
# ------------------------------------------------------------------
DATA_DIR = r""   # folder that holds the project's data, e.g. r"C:/myproject/data"
OUT_DIR  = r""   # folder where data_inventory.csv is saved -- somewhere the AI agent
                 # CAN read, outside the data folder, e.g. r"C:/myproject-code"
INCLUDE_VARNAMES = True   # set to False if variable names themselves are sensitive

# ------------------------------------------------------------------
# 2. Setup
# ------------------------------------------------------------------
import csv
import datetime
import hashlib
from pathlib import Path

if not DATA_DIR or not OUT_DIR:
    raise SystemExit("Set DATA_DIR and OUT_DIR at the top of the script.")

data_dir = Path(DATA_DIR).resolve()
if not data_dir.is_dir():
    raise SystemExit(f"Data folder not found: {data_dir}")

DATA_EXTS = {
    "dta", "sav", "csv", "tsv", "txt", "xlsx", "xls", "rds", "rdata", "rda",
    "parquet", "feather", "json", "shp", "dbf", "gpkg", "geojson", "tif", "zip", "gz",
}

# ------------------------------------------------------------------
# 3. Structure only: variable names and sheet names, never values
# ------------------------------------------------------------------
missing_packages = set()


def sha256(path):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def describe_file(path, ext):
    """Return (variable names or None, sheet names)."""
    sheets = []
    try:
        if ext in ("csv", "tsv"):
            with open(path, newline="", encoding="utf-8-sig", errors="replace") as fh:
                header = next(csv.reader(fh, delimiter="\t" if ext == "tsv" else ","), [])
            return header, sheets

        if ext == "dta":
            import pandas as pd
            # reads one row to get the column names; nothing is kept
            with pd.read_stata(path, iterator=True, chunksize=1) as reader:
                return list(next(iter(reader)).columns), sheets

        if ext == "sav":
            import pyreadstat
            _, meta = pyreadstat.read_sav(str(path), metadataonly=True)
            return list(meta.column_names), sheets

        if ext == "xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True)
            sheets = wb.sheetnames
            first_row = next(wb[sheets[0]].iter_rows(max_row=1, values_only=True), ())
            wb.close()
            return [str(c) for c in first_row if c is not None], sheets

        if ext == "xls":
            import pandas as pd
            xls = pd.ExcelFile(path)
            sheets = xls.sheet_names
            return list(xls.parse(sheets[0], nrows=0).columns), sheets

        if ext == "parquet":
            import pyarrow.parquet as pq
            return list(pq.read_schema(path).names), sheets

    except ImportError as err:
        missing_packages.add(err.name)
    except Exception:
        pass
    return None, sheets


# ------------------------------------------------------------------
# 4. Build and save the report
# ------------------------------------------------------------------
files = sorted(
    p for p in data_dir.rglob("*")
    if p.is_file() and p.suffix.lower().lstrip(".") in DATA_EXTS
)

today = datetime.date.today().isoformat()
rows = []
for p in files:
    ext = p.suffix.lower().lstrip(".")
    variables, sheets = describe_file(p, ext)
    stat = p.stat()
    rows.append({
        "filename": p.name,
        "path": p.relative_to(data_dir).as_posix(),
        "sha256sum": sha256(p),
        "date": today,
        "modified": datetime.datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S"),
        "extension": ext,
        "size_bytes": stat.st_size,
        "n_vars": "" if variables is None else len(variables),
        "sheets": " | ".join(sheets),
        "variables": " ".join(variables) if (variables and INCLUDE_VARNAMES) else "",
        # --- to fill in by the author (external inputs only) ---
        "source": "",
        "url": "",
        "access_date": "",
        "license": "",
        "availability": "",   # public / restricted / forthcoming
        "notes": "",
    })

fields = ["filename", "path", "sha256sum", "date", "modified", "extension", "size_bytes",
          "n_vars", "sheets", "variables",
          "source", "url", "access_date", "license", "availability", "notes"]

out_file = Path(OUT_DIR) / "data_inventory.csv"
with open(out_file, "w", newline="", encoding="utf-8") as fh:
    writer = csv.DictWriter(fh, fieldnames=fields)
    writer.writeheader()
    writer.writerows(rows)

print(f"Data inventory saved to: {out_file} ({len(rows)} files)")
if missing_packages:
    print("Variable names were skipped for some files because these packages are not "
          f"installed: {', '.join(sorted(missing_packages))}")
print("Review it before sharing: it should contain no data values.")
